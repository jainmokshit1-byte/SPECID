"""Upload, mapping, ingest and quality report (TRD TR-MOD-20; PRD FR-101-106, FR-205, FR-1005).

- `parse_table`: CSV (encoding detection: UTF-8, then charset-normalizer on the first 1 MB, then
  cp1252, latin-1) or XLSX (first sheet, read-only).
- `suggest_mapping`: header synonyms of TRD Appendix H, including the SAP preset.
- `ingest_batch`: apply the saved mapping, harmonise units of measure, write records and specs,
  compute the quality report and the health score (DEC-35).
- `ingest_procurement`: purchase lines of one batch, vendor hashed with the CPSE salt (TR-MOD-21).

Services never import FastAPI; they raise typed errors (`app.services.errors`).
"""

import csv
import hashlib
import io
import re
import uuid
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from charset_normalizer import from_bytes
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.core.templates import Template
from app.core.types import Dictionary
from app.core.units import uom_canonical
from app.db.copy import copy_rows
from app.db.models import AppUser, Cpse, UploadBatch
from app.schemas.jsonb import BatchQuality, ColumnMapping, RawRow
from app.services import audit
from app.services.errors import Conflict, Forbidden, Invalid, NotFound
from app.services.specs import RecordText, build_spec, record_text, spec_row, store_specs

MAX_BYTES = 50 * 1024 * 1024  # FR-101
MAX_ROWS = 200_000
SAMPLE_ROWS = 5
FIELDS = (
    "legacy_code short_text long_text uom mat_group manufacturer mpn plant criticality "
    "annual_value annual_qty stock_qty"
).split()
TEXT_FIELDS = ("short_text", "long_text")
NUMERIC = ("annual_value", "annual_qty", "stock_qty")
SAP_MARKERS = frozenset({"matnr", "maktx", "meins", "matkl", "werks", "mfrnr", "mfrpn"})
PROCUREMENT_HEADERS = {
    "legacy_code": ["legacy code", "material code", "matnr", "item code", "code"],
    "po_date": ["po date", "date", "order date", "bedat"],
    "qty": ["qty", "quantity", "menge"],
    "uom": ["uom", "unit", "meins"],
    "unit_price": ["unit price", "price", "rate", "netpr"],
    "currency": ["currency", "waers"],
    "vendor": ["vendor", "supplier", "lifnr"],
    "plant": ["plant", "werks"],
}
_NON_ALNUM = re.compile(r"[^a-z0-9]+")
_NUM = re.compile(r"[^0-9.\-]")


def norm_header(header: str) -> str:
    """Lower-case, punctuation ignored: `MARA-MATNR` and `Mara Matnr` compare equal."""
    return _NON_ALNUM.sub(" ", header.lower()).strip()


# ---------------------------------------------------------------- parsing
@dataclass(frozen=True)
class ParsedTable:
    encoding: str
    headers: list[str]
    rows: list[list[str]]


def _decode(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig",):
        try:
            return data.decode(enc), "utf-8"
        except UnicodeDecodeError:
            pass
    best = from_bytes(data[: 1024 * 1024]).best()
    tried = [best.encoding] if best and best.encoding else []
    for enc in [*tried, "cp1252", "latin-1"]:
        try:
            return data.decode(enc), enc
        except (UnicodeDecodeError, LookupError):
            continue
    raise Invalid("The file text encoding could not be read.", title="Unreadable file")


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse_table(data: bytes, filename: str) -> ParsedTable:
    if not data:
        raise Invalid("The file is empty.", title="Empty file")
    if len(data) > MAX_BYTES:
        raise Invalid(
            f"Files are limited to {MAX_BYTES // (1024 * 1024)} MB.", title="File too big"
        )
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            grid = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        except Exception as exc:  # corrupt workbook
            raise Invalid("The Excel file could not be read.", title="Unreadable file") from exc
        encoding = "xlsx"
    elif name.endswith((".csv", ".txt", ".tsv")):
        content, encoding = _decode(data)
        sample = content[:20_000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        grid = [[c.strip() for c in row] for row in csv.reader(io.StringIO(content), dialect)]
    else:
        raise Invalid("Upload a .csv or .xlsx file.", title="Unsupported file type")
    grid = [row for row in grid if any(c for c in row)]
    if not grid:
        raise Invalid("The file has no rows.", title="Empty file")
    headers = [h or f"column {i + 1}" for i, h in enumerate(grid[0])]
    if len(set(norm_header(h) for h in headers)) != len(headers):
        raise Invalid("Two columns have the same header.", title="Duplicate column headers")
    rows = [row + [""] * (len(headers) - len(row)) for row in grid[1:]]
    if len(rows) > MAX_ROWS:
        raise Invalid(f"Files are limited to {MAX_ROWS:,} rows.", title="Too many rows")
    return ParsedTable(encoding, headers, rows)


# ---------------------------------------------------------------- mapping
def suggest_mapping(
    headers: Sequence[str], dictionary: Dictionary
) -> tuple[dict[str, str], str | None]:
    """({source header: target field}, preset). Each target gets at most one source header."""
    lookup: dict[str, str] = {}
    for target, synonyms in dictionary.header_synonyms.items():
        for word in (target.replace("_", " "), *synonyms):
            lookup.setdefault(norm_header(word), target)
    mapping: dict[str, str] = {}
    taken: set[str] = set()
    for header in headers:
        target = lookup.get(norm_header(header))
        if target and target not in taken and target in FIELDS:
            mapping[header] = target
            taken.add(target)
    parts = {w for h in headers for w in norm_header(h).split()}
    preset = "SAP" if len(SAP_MARKERS & parts) >= 3 else None
    return mapping, preset


def validate_mapping(mapping: Mapping[str, str], headers: Sequence[str]) -> dict[str, str]:
    clean = ColumnMapping.model_validate(dict(mapping)).root
    unknown = [h for h in clean if h not in headers]
    if unknown:
        raise Invalid(
            f"These columns are not in the file: {', '.join(unknown)}.", title="Bad mapping"
        )
    bad = sorted({t for t in clean.values() if t not in FIELDS})
    if bad:
        raise Invalid(f"Unknown target field(s): {', '.join(bad)}.", title="Bad mapping")
    dup = [t for t, n in Counter(clean.values()).items() if n > 1]
    if dup:
        raise Invalid(f"Two columns map to {', '.join(dup)}.", title="Bad mapping")
    targets = set(clean.values())
    if "legacy_code" not in targets:
        raise Invalid("Map a column to legacy_code.", title="Bad mapping")
    if not targets & set(TEXT_FIELDS):
        raise Invalid("Map a column to short_text or long_text.", title="Bad mapping")
    return clean


# ---------------------------------------------------------------- batches
def upload_path(upload_dir: str, batch_id: uuid.UUID) -> Path:
    return Path(upload_dir) / f"{batch_id}.bin"


def _scope_check(actor: AppUser, cpse: Cpse) -> None:
    """MAKER and CHECKER act for their own CPSE only (PRD section 2); ADMIN for any."""
    if actor.role in ("MAKER", "CHECKER") and actor.cpse_id != cpse.id:
        raise Forbidden("You can upload only for your own CPSE.", title="No access")


def resolve_cpse(session: Session, actor: AppUser, code: str | None) -> Cpse:
    if actor.role in ("MAKER", "CHECKER"):
        cpse = session.get(Cpse, actor.cpse_id) if actor.cpse_id else None
        if cpse is None:
            raise Invalid("Your account has no CPSE.", title="No CPSE")
        if code and code != cpse.code:
            raise Forbidden("You can upload only for your own CPSE.", title="No access")
        return cpse
    if not code:
        raise Invalid("Say which CPSE this file belongs to.", title="CPSE required")
    cpse = session.scalar(select(Cpse).where(Cpse.code == code))
    if cpse is None:
        raise NotFound(f"No CPSE with code {code}.")
    return cpse


def get_batch(session: Session, batch_id: uuid.UUID) -> UploadBatch:
    batch = session.get(UploadBatch, batch_id)
    if batch is None:
        raise NotFound("No such batch.")
    return batch


@dataclass(frozen=True)
class UploadResult:
    batch: UploadBatch
    table: ParsedTable | None  # None when the file was already ingested
    suggested: dict[str, str]
    preset: str | None
    already_ingested: bool


def create_batch(
    session: Session,
    actor: AppUser,
    *,
    cpse: Cpse,
    filename: str,
    data: bytes,
    is_synthetic: bool,
    dictionary: Dictionary,
    upload_dir: str,
) -> UploadResult:
    _scope_check(actor, cpse)
    sha = hashlib.sha256(data).hexdigest()
    same = session.scalar(
        select(UploadBatch).where(
            UploadBatch.cpse_id == cpse.id,
            UploadBatch.file_sha256 == sha,
            UploadBatch.status == "INGESTED",
        )
    )
    if same is not None:  # FR-105: re-uploading a file creates nothing new
        return UploadResult(same, None, {}, None, True)
    table = parse_table(data, filename)
    suggested, preset = suggest_mapping(table.headers, dictionary)
    batch = UploadBatch(
        cpse_id=cpse.id,
        filename=Path(filename).name[:255],
        file_sha256=sha,
        status="UPLOADED",
        column_mapping=suggested or None,
        row_count=len(table.rows),
        is_synthetic=is_synthetic,
        created_by=actor.id,
    )
    session.add(batch)
    session.flush()
    path = upload_path(upload_dir, batch.id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    audit.record(
        session,
        actor_id=actor.id,
        action="BATCH_UPLOADED",
        object_type="upload_batch",
        object_id=str(batch.id),
        after={"cpse": cpse.code, "filename": batch.filename, "rows": len(table.rows),
               "sha256": sha, "is_synthetic": is_synthetic},
    )  # fmt: skip
    return UploadResult(batch, table, suggested, preset, False)


def load_table(batch: UploadBatch, upload_dir: str) -> ParsedTable:
    path = upload_path(upload_dir, batch.id)
    if not path.exists():
        raise Conflict("The uploaded file is no longer on the server. Upload it again.")
    return parse_table(path.read_bytes(), batch.filename)


def save_mapping(
    session: Session,
    actor: AppUser,
    batch: UploadBatch,
    mapping: Mapping[str, str],
    upload_dir: str,
) -> UploadBatch:
    cpse = session.get(Cpse, batch.cpse_id)
    assert cpse is not None
    _scope_check(actor, cpse)
    if batch.status not in ("UPLOADED", "MAPPED"):
        raise Conflict("This batch is already ingested; its mapping cannot change.")
    clean = validate_mapping(mapping, load_table(batch, upload_dir).headers)
    batch.column_mapping = clean
    batch.status = "MAPPED"
    audit.record(
        session,
        actor_id=actor.id,
        action="BATCH_MAPPED",
        object_type="upload_batch",
        object_id=str(batch.id),
        after={"mapping": clean},
    )
    return batch


# ---------------------------------------------------------------- values
def _num(value: str) -> Decimal | None:
    cleaned = _NUM.sub("", value.replace(",", ""))
    if cleaned in ("", "-", "."):
        return None
    try:
        d = Decimal(cleaned)
    except InvalidOperation:
        return None
    return d if d >= 0 else None


def _criticality(value: str) -> str | None:
    v = value.strip().upper()
    return {
        "Y": "Y",
        "YES": "Y",
        "TRUE": "Y",
        "1": "Y",
        "N": "N",
        "NO": "N",
        "FALSE": "N",
        "0": "N",
    }.get(v)


def _content_hash(fields: Mapping[str, str]) -> str:
    canon = "\x1f".join(f"{k}={fields.get(k, '')}" for k in FIELDS)
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class MappedRow:
    fields: dict[str, str]
    raw: dict[str, str]


def map_rows(table: ParsedTable, mapping: Mapping[str, str]) -> Iterable[MappedRow]:
    index = {header: i for i, header in enumerate(table.headers)}
    for row in table.rows:
        fields = {t: row[index[h]].strip() for h, t in mapping.items()}
        yield MappedRow(fields, dict(zip(table.headers, row, strict=True)))


# ---------------------------------------------------------------- ingest
def _ratio(n: int | float, total: int | float) -> float:
    return round(n / total, 4) if total else 0.0


def ingest_batch(
    session: Session,
    actor: AppUser,
    batch: UploadBatch,
    *,
    dictionary: Dictionary,
    templates: Mapping[str, Template],
    threshold: float,
    upload_dir: str,
    model: Any = None,
) -> BatchQuality:
    """Ingest the saved mapping (idempotent: an ingested batch returns its quality again)."""
    cpse = session.get(Cpse, batch.cpse_id)
    assert cpse is not None
    _scope_check(actor, cpse)
    if batch.status == "INGESTED" and batch.quality:
        return BatchQuality.model_validate(batch.quality)
    if not batch.column_mapping:
        raise Invalid("Map the columns first.", title="No mapping")
    table = load_table(batch, upload_dir)
    mapping = validate_mapping(batch.column_mapping, table.headers)
    batch.status = "INGESTING"

    known = {  # (legacy_code, content_hash) already stored for this CPSE: unchanged rows skip
        (code, h)
        for code, h in session.execute(
            text("SELECT legacy_code, content_hash FROM material_record WHERE cpse_id = :c"),
            {"c": cpse.id},
        )
    }
    seen: set[str] = set()
    records: list[tuple[Any, ...]] = []
    texts: list[RecordText] = []
    present = Counter[str]()
    duplicate_codes = rejected = skipped = empty_short = over_40 = ambiguous = uom_unknown = 0
    for mr in map_rows(table, mapping):
        f = mr.fields
        code = f.get("legacy_code", "")
        short, long = f.get("short_text", ""), f.get("long_text", "")
        if not code or not (short or long):
            rejected += 1
            continue
        if code in seen:
            duplicate_codes += 1
            continue
        seen.add(code)
        chash = _content_hash(f)
        if (code, chash) in known:
            skipped += 1
            continue
        short = short or long[:40].rstrip()
        empty_short += not f.get("short_text")
        over_40 += len(f.get("short_text", "")) > 40
        raw_uom = f.get("uom", "")
        canon, amb = uom_canonical(raw_uom, dictionary) if raw_uom else (None, False)
        ambiguous += amb
        uom_unknown += bool(raw_uom) and canon is None and not amb
        uom_unknown += not raw_uom
        rid = uuid.uuid4()
        for name in FIELDS:
            present[name] += bool(f.get(name))
        records.append(
            (
                rid, batch.id, cpse.id, code, short, long or None, raw_uom or None, canon,
                f.get("mat_group") or None, f.get("manufacturer") or None, f.get("mpn") or None,
                f.get("plant") or None, _criticality(f.get("criticality", "")),
                _num(f.get("annual_value", "")), _num(f.get("annual_qty", "")),
                _num(f.get("stock_qty", "")), chash,
                RawRow.model_validate(mr.raw).root,
            )
        )  # fmt: skip
        texts.append(RecordText(rid, record_text(short, long), f.get("mpn") or None,
                                f.get("manufacturer") or None))  # fmt: skip

    conn = session.connection()
    copy_rows(
        conn, "material_record",
        ["id", "batch_id", "cpse_id", "legacy_code", "short_text", "long_text", "uom",
         "uom_canonical", "mat_group", "manufacturer", "mpn", "plant", "criticality",
         "annual_value", "annual_qty", "stock_qty", "content_hash", "raw"],
        records,
    )  # fmt: skip

    categories = Counter[str]()
    parsed = Counter[str]()
    completeness: list[float] = []
    rows_for_specs = []
    for rec in texts:
        spec = build_spec(rec, dictionary, threshold, model)
        rows_for_specs.append(spec_row(rec, spec, templates, dictionary))
        categories[spec.category or "UNRECOGNISED"] += 1
        t = templates.get(spec.category or "")
        if t is not None:
            full = all(spec.attrs.get(c) is not None for c in t.core)
            parsed[t.category] += full
            completeness.append(float(rows_for_specs[-1][10] or 0))
    store_specs(conn, rows_for_specs)

    n = len(records)
    recognised = n - categories.get("UNRECOGNISED", 0)
    components = {
        "descriptions": _ratio(n - empty_short, n),
        "unique_codes": _ratio(n, n + duplicate_codes),
        "recognised_category": _ratio(recognised, n),
        "spec_completeness": (
            round(sum(completeness) / len(completeness), 4) if completeness else 0.0
        ),
        "uom_clean": _ratio(n - uom_unknown - ambiguous, n),
    }
    quality = BatchQuality(
        rows=n,
        empty_short_text=empty_short,
        short_text_over_40=over_40,
        duplicate_legacy_codes=duplicate_codes,
        completeness={name: _ratio(present[name], n) for name in FIELDS},
        category_share={c: _ratio(k, n) for c, k in sorted(categories.items())},
        core_parse_rate={c: _ratio(parsed[c], categories[c]) for c in sorted(parsed)},
        uom_ambiguous=ambiguous,
        rejected_rows=rejected,
        uom_unknown=uom_unknown,
        skipped_unchanged=skipped,
        health_score=round(100 * sum(components.values()) / len(components)) if n else None,
        health_components=components,
    )
    batch.quality = quality.model_dump(mode="json")
    batch.row_count = n
    batch.status = "INGESTED"
    audit.record(
        session,
        actor_id=actor.id,
        action="BATCH_INGESTED",
        object_type="upload_batch",
        object_id=str(batch.id),
        after={"cpse": cpse.code, "records": n, "rejected": rejected, "skipped": skipped,
               "health_score": quality.health_score},
    )  # fmt: skip
    return quality


# ---------------------------------------------------------------- procurement
def _vendor_hash(salt: str, vendor: str) -> str | None:
    vendor = vendor.strip()
    return hashlib.sha256((salt + vendor).encode("utf-8")).hexdigest() if vendor else None


def _date(value: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    return None


def ingest_procurement(
    session: Session, actor: AppUser, batch: UploadBatch, *, filename: str, data: bytes
) -> dict[str, int]:
    """Purchase lines for the records of an ingested batch (FR-107). Returns counts."""
    cpse = session.get(Cpse, batch.cpse_id)
    assert cpse is not None
    _scope_check(actor, cpse)
    if batch.status != "INGESTED":
        raise Conflict("Ingest the material file before its purchase history.")
    table = parse_table(data, filename)
    lookup = {norm_header(w): t for t, words in PROCUREMENT_HEADERS.items() for w in (t, *words)}
    cols: dict[str, int] = {}
    for i, header in enumerate(table.headers):
        target = lookup.get(norm_header(header))
        if target and target not in cols:
            cols[target] = i
    missing = [c for c in ("legacy_code", "po_date", "qty") if c not in cols]
    if missing:
        raise Invalid(f"Missing column(s): {', '.join(missing)}.", title="Bad procurement file")

    by_code = {
        code: rid
        for rid, code in session.execute(
            text("SELECT id, legacy_code FROM material_record WHERE batch_id = :b"), {"b": batch.id}
        )
    }
    existing = {
        (r[0], r[1], r[2], r[3], r[4])
        for r in session.execute(
            text(
                "SELECT record_id, po_date, qty, unit_price, vendor_hash FROM procurement_line"
                " WHERE record_id = ANY(:ids)"
            ),
            {"ids": list(by_code.values())},
        )
    }
    salt = cpse.vendor_salt
    lines: list[tuple[Any, ...]] = []
    unmatched = bad = duplicates = 0
    for row in table.rows:

        def get(name: str, row: list[str] = row) -> str:
            return row[cols[name]].strip() if name in cols else ""

        rid = by_code.get(get("legacy_code"))
        po, qty = _date(get("po_date")), _num(get("qty"))
        if rid is None:
            unmatched += 1
            continue
        if po is None or qty is None:
            bad += 1
            continue
        price = _num(get("unit_price"))
        vh = _vendor_hash(salt, get("vendor"))
        key = (rid, po, qty, price, vh)
        if key in existing:
            duplicates += 1
            continue
        existing.add(key)
        lines.append((rid, po, qty, get("uom") or None, price, get("currency") or "INR", vh,
                      get("plant") or None))  # fmt: skip
    conn = session.connection()
    copy_rows(
        conn, "procurement_line",
        ["record_id", "po_date", "qty", "uom", "unit_price", "currency", "vendor_hash", "plant"],
        lines,
    )  # fmt: skip
    # annual figures from the last 12 months of the data, only where the master file gave none
    conn.execute(
        text("""
            WITH asof AS (
              SELECT max(p.po_date) AS d FROM procurement_line p
              JOIN material_record m ON m.id = p.record_id WHERE m.batch_id = :b
            ), agg AS (
              SELECT p.record_id, sum(p.qty) AS q, sum(p.qty * COALESCE(p.unit_price, 0)) AS v
              FROM procurement_line p JOIN material_record m ON m.id = p.record_id, asof
              WHERE m.batch_id = :b AND p.po_date > asof.d - interval '12 months'
              GROUP BY p.record_id
            )
            UPDATE material_record m
            SET annual_qty = COALESCE(m.annual_qty, agg.q),
                annual_value = COALESCE(m.annual_value, round(agg.v))
            FROM agg WHERE m.id = agg.record_id
            """),
        {"b": batch.id},
    )
    counts = {"lines": len(lines), "unmatched": unmatched, "invalid": bad, "duplicates": duplicates}
    audit.record(
        session,
        actor_id=actor.id,
        action="PROCUREMENT_INGESTED",
        object_type="upload_batch",
        object_id=str(batch.id),
        after={"cpse": cpse.code, **counts},
    )
    return counts
